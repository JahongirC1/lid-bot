"""
Lid -> Excel bot + Guruh qidiruvi dasturi
-----------------------------------------
1) Lid xabarini (🧾 Nomi: ... 📝 Ma'lumotlar: ...) botga tashlaysiz,
   bot uni ajratib oladi va Excel faylga qator qilib qo'shadi.
2) Bot manzilida (Render linki) "Guruh qidiruvi" dasturi ochiladi (index.html).
   Dastur lidlarni botdan oladi, ularni mos guruhga avtomatik joylaydi,
   siz holatini belgilaysiz - hammasi Excel'ga ham yoziladi.
   Dasturga kirish linki: botga /dastur yozing.

Env o'zgaruvchilar (Render > Environment):
  BOT_TOKEN   - BotFather bergan token (majburiy)
  ADMIN_IDS   - bosh admin(lar) ID si, vergul bilan: 123456789,987654321
                Yordamchilar botning o'zida qo'shiladi: yordamchi /start yozadi,
                bosh adminga "Ruxsat berish" tugmasi keladi. Ro'yxat: /yordamchilar
  SELF_URL    - (ixtiyoriy) bot linki; bo'lmasa Render bergan RENDER_EXTERNAL_URL olinadi
  SEND_DELAY  - oxirgi lid tushgandan keyin necha soniyada Excel yuborilsin (standart 4)
"""

import asyncio
import hashlib
import hmac
import io
import json
import logging
import os
import re
import threading
import time
import urllib.request
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, InputMediaDocument, Update
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

logging.basicConfig(format="%(asctime)s %(levelname)s %(message)s", level=logging.INFO)
log = logging.getLogger("lidbot")
logging.getLogger("httpx").setLevel(logging.WARNING)  # token loglarda chiqmasin

BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
ADMIN_IDS = {int(x) for x in re.findall(r"\d+", os.environ.get("ADMIN_IDS", ""))}
SELF_URL = (os.environ.get("SELF_URL") or os.environ.get("RENDER_EXTERNAL_URL", "")).rstrip("/")
SEND_DELAY = float(os.environ.get("SEND_DELAY", "4"))
LOCAL_FILE = "lidlar.xlsx"
DATE_FMT = "%d-%m-%Y %H:%M:%S"
HERE = os.path.dirname(os.path.abspath(__file__))
INDEX_FILE = os.path.join(HERE, "index.html")

# Dasturga kirish kaliti - bot tokenidan yasaladi (alohida parol kerak emas)
WEB_KEY = hashlib.sha256(("dastur:" + BOT_TOKEN).encode()).hexdigest()[:20]

HOLATLAR = ["Yangi", "Qo'ng'iroq qilindi", "Guruhga qo'shildi", "Rad etdi"]

# ---------------------------------------------------------------- Excel tuzilishi
# (kalit, sarlavha, kenglik)
INFO_COLS = [
    ("daraja", "Daraja", 10),
    ("kurs", "Kurs", 15),
    ("kun", "Kunlar", 11),
    ("vaqt", "Vaqt", 14),
    ("ism", "Ism-familiya", 24),
    ("tel", "Telefon", 17),
]
LEFT_COLS = [("n", "№", 6)]
RIGHT_COLS = [
    ("guruh", "Guruh", 48),
    ("holat", "Holat", 19),
    ("manba", "Manba", 13),
    ("sana", "Sana", 19),
]
ALL_COLS = LEFT_COLS + INFO_COLS + RIGHT_COLS

rows: list[dict] = []  # barcha lidlar (xotirada)
lock = threading.RLock()  # rows ga bot va dastur bir vaqtda tegmasligi uchun
state = {"dirty": False}  # o'zgarish bo'ldi - hammaning pin qilingan Excel'i yangilansin
helpers: dict[int, str] = {}  # yordamchilar: {telegram_id: ism} - Excel'ning 2-varag'ida saqlanadi
requests: dict[int, str] = {}  # ruxsat so'raganlar
dead: set[int] = set()  # botga hali /start yozmagan (xabar yuborib bo'lmaydigan) chatlar

# ---------------------------------------------------------------- Xabarni o'qish
QUESTION_RE = re.compile(r"^\s*\d+\s*[.)]\s*(.+?)\s*:\s*(.*)$")
IGNORED = ("form name", "ad name", "campaign", "adset", "ad set", "forma", "e'lon")


def field_for(question: str) -> str | None:
    q = question.lower().strip()
    if any(w in q for w in IGNORED):
        return None
    if "daraja" in q or "level" in q:
        return "daraja"
    if "kurs" in q or "course" in q:
        return "kurs"
    if "vaqt" in q or "time" in q or "soat" in q:
        return "vaqt"
    if "kun" in q or "day" in q:
        return "kun"
    if "phone" in q or "telefon" in q or "raqam" in q or "nomer" in q:
        return "tel"
    if "full name" in q or "ism" in q or "fio" in q or "f.i.sh" in q or q == "name":
        return "ism"
    return None


def clean_phone(p: str) -> str:
    digits = re.sub(r"\D", "", p)
    if len(digits) == 9:  # 900341519 -> +998900341519
        digits = "998" + digits
    return "+" + digits if digits else p.strip()


def parse_block(lines: list[str]) -> dict | None:
    lead: dict = {}
    for line in lines:
        m = QUESTION_RE.match(line)
        if m:
            key = field_for(m.group(1))
            if key and m.group(2).strip() and key not in lead:
                lead[key] = m.group(2).strip()
            continue
        for key, label in (("manba", "Manba:"), ("sana", "Sana:")):
            if label in line:
                lead[key] = line.split(label, 1)[1].strip()
    if not (lead.get("tel") or lead.get("ism")):
        return None
    if lead.get("tel"):
        lead["tel"] = clean_phone(lead["tel"])
    if lead.get("sana"):
        try:
            lead["sana"] = datetime.strptime(lead["sana"], DATE_FMT)
        except ValueError:
            pass
    lead["holat"] = "Yangi"
    return lead


def parse_message(text: str) -> list[dict]:
    """Bitta xabarda bir nechta lid bo'lsa ham ajratadi ('Nomi:' qatoridan boshlab)."""
    blocks, cur = [], []
    for line in text.splitlines():
        if "Nomi:" in line and cur:
            blocks.append(cur)
            cur = []
        cur.append(line)
    if cur:
        blocks.append(cur)
    return [l for l in (parse_block(b) for b in blocks) if l]


def lead_key(lead: dict) -> tuple:
    sana = lead.get("sana")
    if isinstance(sana, datetime):
        sana = sana.strftime(DATE_FMT)
    return (re.sub(r"\D", "", str(lead.get("tel", ""))), str(sana or ""), str(lead.get("ism", "")).lower())


def lead_id(lead: dict) -> str:
    return hashlib.sha1("|".join(lead_key(lead)).encode()).hexdigest()[:12]


def add_leads(new: list[dict]) -> tuple[int, int]:
    with lock:
        known = {lead_key(r) for r in rows}
        added = dup = 0
        for lead in new:
            k = lead_key(lead)
            if k in known:
                dup += 1
                continue
            known.add(k)
            lead.setdefault("holat", "Yangi")
            rows.append(lead)
            added += 1
        return added, dup


# ---------------------------------------------------------------- Excel yasash
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="2E75B6")
ZEBRA = PatternFill("solid", fgColor="EAF1FB")
WHITE_BOLD = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="9BB7D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HOLAT_FONT = {
    "Guruhga qo'shildi": Font(bold=True, color="0D6B55"),
    "Rad etdi": Font(color="B42318"),
    "Qo'ng'iroq qilindi": Font(color="8F540B"),
}


def build_workbook() -> bytes:
    with lock:
        data_rows = [dict(r) for r in rows]
    wb = Workbook()
    ws = wb.active
    ws.title = "Lidlar"

    info_start = len(LEFT_COLS) + 1
    info_end = info_start + len(INFO_COLS) - 1

    # 1-qator: "Ma'lumotlar" umumiy sarlavha, qolganlari 2 qatorga birlashtirilgan
    for i, (_, title, width) in enumerate(ALL_COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if not (info_start <= i <= info_end):
            ws.cell(row=1, column=i, value=title)
            ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
        else:
            ws.cell(row=2, column=i, value=title)
    ws.cell(row=1, column=info_start, value="Ma'lumotlar")
    ws.merge_cells(start_row=1, start_column=info_start, end_row=1, end_column=info_end)

    for r in (1, 2):
        for c in range(1, len(ALL_COLS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = SUB_FILL if (r == 2 and info_start <= c <= info_end) else HEAD_FILL
            cell.font = WHITE_BOLD
            cell.alignment = CENTER
            cell.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20

    for idx, lead in enumerate(data_rows, start=1):
        r = idx + 2
        for c, (key, _, _) in enumerate(ALL_COLS, start=1):
            val = idx if key == "n" else lead.get(key, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = BORDER
            cell.alignment = LEFT if key in ("kurs", "ism", "guruh") else CENTER
            if key == "tel":
                cell.number_format = "@"
            if key == "sana" and isinstance(val, datetime):
                cell.number_format = "dd.mm.yyyy hh:mm"
            if key == "holat" and val in HOLAT_FONT:
                cell.font = HOLAT_FONT[val]
            if idx % 2 == 0:
                cell.fill = ZEBRA

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(ALL_COLS))}{max(len(data_rows) + 2, 2)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"

    hs = wb.create_sheet("Yordamchilar")
    hs.append(["Telegram ID", "Ism"])
    for c in (1, 2):
        hs.cell(row=1, column=c).font = WHITE_BOLD
        hs.cell(row=1, column=c).fill = HEAD_FILL
    hs.column_dimensions["A"].width = 16
    hs.column_dimensions["B"].width = 30
    for uid, name in sorted(helpers.items()):
        hs.append([str(uid), name])
    wb.active = 0

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    try:
        with open(LOCAL_FILE, "wb") as f:
            f.write(data)
    except OSError:
        pass
    return data


def load_workbook_bytes(data: bytes) -> list[dict]:
    """Bot yasagan Excel'ni qayta o'qiydi (tiklash uchun) - ustunlarni sarlavha nomidan taniydi."""
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    if "Yordamchilar" in wb.sheetnames:
        for row in wb["Yordamchilar"].iter_rows(min_row=2, values_only=True):
            if row and row[0] and str(row[0]).strip().isdigit():
                helpers.setdefault(int(str(row[0]).strip()), str(row[1] or ""))
    by_title = {title: key for key, title, _ in ALL_COLS if key != "n"}
    colmap = {}
    for c in range(1, ws.max_column + 1):
        title = ws.cell(row=2, column=c).value or ws.cell(row=1, column=c).value
        if title in by_title:
            colmap[c] = by_title[title]
    loaded = []
    for r in range(3, ws.max_row + 1):
        lead = {}
        for c, key in colmap.items():
            val = ws.cell(row=r, column=c).value
            if val not in (None, ""):
                lead[key] = val if isinstance(val, datetime) else str(val)
        if lead.get("tel") or lead.get("ism"):
            lead.setdefault("holat", "Yangi")
            loaded.append(lead)
    return loaded


# ---------------------------------------------------------------- Dastur uchun API
def lead_json(lead: dict) -> dict:
    sana = lead.get("sana")
    return {
        "id": lead_id(lead),
        "ism": str(lead.get("ism", "")),
        "tel": str(lead.get("tel", "")),
        "daraja": str(lead.get("daraja", "")),
        "kurs": str(lead.get("kurs", "")),
        "kun": str(lead.get("kun", "")),
        "vaqt": str(lead.get("vaqt", "")),
        "manba": str(lead.get("manba", "")),
        "sana": sana.strftime("%Y-%m-%dT%H:%M:%S") if isinstance(sana, datetime) else str(sana or ""),
        "holat": str(lead.get("holat", "Yangi")),
        "guruh": str(lead.get("guruh", "")),
    }


def api_update(body: dict) -> dict | None:
    with lock:
        for lead in rows:
            if lead_id(lead) == body.get("id"):
                if "holat" in body and body["holat"] in HOLATLAR:
                    lead["holat"] = body["holat"]
                if "guruh" in body:
                    lead["guruh"] = str(body["guruh"])[:200]
                state["dirty"] = True
                return lead_json(lead)
    return None


def dastur_link() -> str:
    return f"{SELF_URL}/?k={WEB_KEY}" if SELF_URL else ""


class Web(BaseHTTPRequestHandler):
    def _send(self, code: int, body: bytes, ctype: str) -> None:
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        if self.command != "HEAD":
            self.wfile.write(body)

    def _json(self, code: int, obj) -> None:
        self._send(code, json.dumps(obj, ensure_ascii=False).encode(), "application/json; charset=utf-8")

    def _authed(self) -> bool:
        k = self.headers.get("X-Key") or parse_qs(urlparse(self.path).query).get("k", [""])[0]
        return hmac.compare_digest(k, WEB_KEY)

    def do_HEAD(self):
        self.do_GET()

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/leads":
            if not self._authed():
                return self._json(403, {"error": "kalit noto'g'ri"})
            with lock:
                data = [lead_json(r) for r in rows]
            return self._json(200, {"leads": data, "holatlar": HOLATLAR})
        if path in ("/", "/index.html"):
            try:
                with open(INDEX_FILE, "rb") as f:
                    return self._send(200, f.read(), "text/html; charset=utf-8")
            except OSError:
                return self._send(200, b"OK", "text/plain")
        return self._send(404, b"Not found", "text/plain")

    def do_POST(self):
        if urlparse(self.path).path != "/api/lead":
            return self._send(404, b"Not found", "text/plain")
        if not self._authed():
            return self._json(403, {"error": "kalit noto'g'ri"})
        try:
            n = int(self.headers.get("Content-Length", "0"))
            body = json.loads(self.rfile.read(min(n, 10000)) or b"{}")
        except (ValueError, json.JSONDecodeError):
            return self._json(400, {"error": "noto'g'ri so'rov"})
        updated = api_update(body)
        if not updated:
            return self._json(404, {"error": "lid topilmadi"})
        return self._json(200, updated)

    def log_message(self, *a):
        pass


def run_web():
    port = int(os.environ.get("PORT", "10000"))
    ThreadingHTTPServer(("0.0.0.0", port), Web).serve_forever()


def self_ping():
    while True:
        time.sleep(9 * 60)
        try:
            urllib.request.urlopen(SELF_URL, timeout=20).read()
        except Exception as e:
            log.warning("Ping xatosi: %s", e)


# ---------------------------------------------------------------- Telegram
pending: dict[int, asyncio.Task] = {}
last_pinned: dict[int, int] = {}


def is_owner(uid: int | None) -> bool:
    return not ADMIN_IDS or uid in ADMIN_IDS


def allowed(update: Update) -> bool:
    u = update.effective_user
    if not u:
        return False
    ok = is_owner(u.id) or u.id in helpers
    if ok:
        dead.discard(u.id)
    return ok


def user_label(u) -> str:
    name = " ".join(x for x in (u.first_name, u.last_name) if x) or "Ismsiz"
    return f"{name} (@{u.username})" if u.username else name


def excel_name() -> str:
    return f"lidlar_{datetime.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"


async def send_excel(chat_id: int, bot, force: bool = False) -> None:
    if not rows and not force:
        await bot.send_message(chat_id, "Hozircha lid yo'q.")
        return
    data = build_workbook()
    msg = await bot.send_document(
        chat_id, document=data, filename=excel_name(), caption=f"📊 Jami: {len(rows)} ta lid"
    )
    # Eng so'nggi faylni pin qilamiz - server qayta ishga tushsa, shundan tiklanadi
    try:
        if chat_id in last_pinned:
            await bot.unpin_chat_message(chat_id, message_id=last_pinned[chat_id])
        await bot.pin_chat_message(chat_id, msg.message_id, disable_notification=True)
        last_pinned[chat_id] = msg.message_id
    except Exception as e:
        log.warning("Pin xatosi: %s", e)


async def persist_loop(app: Application) -> None:
    """O'zgarish bo'lsa, bosh admin va yordamchilarning pin qilingan Excel'ini jimgina yangilaydi."""
    while True:
        await asyncio.sleep(15)
        if not state["dirty"]:
            continue
        state["dirty"] = False
        targets = (set(last_pinned) | set(ADMIN_IDS) | set(helpers)) - dead
        for chat_id in targets:
            try:
                if chat_id in last_pinned:
                    media = InputMediaDocument(
                        media=build_workbook(), filename=excel_name(), caption=f"📊 Jami: {len(rows)} ta lid"
                    )
                    await app.bot.edit_message_media(media=media, chat_id=chat_id, message_id=last_pinned[chat_id])
                else:
                    await send_excel(chat_id, app.bot, force=True)
            except Exception as e:
                err = str(e).lower()
                if "not modified" in err:
                    continue
                if "forbidden" in err or "chat not found" in err or "blocked" in err:
                    dead.add(chat_id)  # botga hali /start yozmagan
                    continue
                log.warning("Excel yangilash xatosi (%s): %s", chat_id, e)
                try:
                    last_pinned.pop(chat_id, None)
                    await send_excel(chat_id, app.bot, force=True)
                except Exception as e2:
                    log.warning("Excel qayta yuborish xatosi: %s", e2)


async def _delayed_send(chat_id: int, bot) -> None:
    try:
        await asyncio.sleep(SEND_DELAY)
    except asyncio.CancelledError:
        return
    await send_excel(chat_id, bot)


def schedule_send(chat_id: int, bot) -> None:
    """Ketma-ket bir nechta lid tashlansa, Excel faqat bir marta - oxirida yuboriladi."""
    t = pending.get(chat_id)
    if t and not t.done():
        t.cancel()
    pending[chat_id] = asyncio.create_task(_delayed_send(chat_id, bot))


async def start(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    u = update.effective_user
    uid = u.id
    if not allowed(update):
        if uid in requests:
            await update.message.reply_text("⏳ So'rovingiz yuborilgan, bosh admin ruxsat berishini kuting.")
            return
        requests[uid] = user_label(u)
        kb = InlineKeyboardMarkup([[
            InlineKeyboardButton("✅ Ruxsat berish", callback_data=f"add:{uid}"),
            InlineKeyboardButton("❌ Rad etish", callback_data=f"no:{uid}"),
        ]])
        sent = False
        for owner in ADMIN_IDS:
            try:
                await ctx.bot.send_message(
                    owner, f"👤 {requests[uid]} (ID: {uid}) botdan yordamchi sifatida foydalanmoqchi.", reply_markup=kb
                )
                sent = True
            except Exception as e:
                log.warning("So'rov yuborilmadi (%s): %s", owner, e)
        await update.message.reply_text(
            "📨 So'rovingiz bosh adminga yuborildi. Ruxsat berilgach, shu yerga xabar keladi."
            if sent else f"Sizga ruxsat yo'q. ID: {uid}"
        )
        return
    role = "bosh admin" if is_owner(uid) else "yordamchi"
    await update.message.reply_text(
        "Assalamu alaykum! Lid xabarlarini shu yerga tashlang (forward yoki nusxa) — "
        "men ularni Excel jadvalga qo'shib boraman.\n\n"
        "/dastur — guruh qidiruvi dasturi (lidlarni guruhga joylash)\n"
        "/excel — jadvalni hozir olish\n"
        "/soni — nechta lid borligi\n"
        "/tozalash — hammasini o'chirish (faqat bosh admin)\n"
        "/yordamchilar — yordamchilar ro'yxati (faqat bosh admin)\n\n"
        "Eski Excel faylni yuborsangiz, undagi lidlarni ham qo'shib oladi.\n\n"
        f"Siz: {role} · ID: {uid}"
    )


async def cmd_dastur(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if not allowed(update):
        return
    link = dastur_link()
    if not link:
        await update.message.reply_text("Bot linki topilmadi. Render'da SELF_URL o'zgaruvchisini qo'ying.")
        return
    await update.message.reply_text(
        f"📋 Guruh qidiruvi dasturi:\n{link}\n\n"
        "Bu link — kalit. Uni boshqalarga bermang (ichida lidlarning raqamlari bor).",
        disable_web_page_preview=True,
    )


def helpers_markup() -> InlineKeyboardMarkup | None:
    if not helpers:
        return None
    return InlineKeyboardMarkup(
        [[InlineKeyboardButton(f"❌ {name or uid}", callback_data=f"del:{uid}")] for uid, name in sorted(helpers.items())]
    )


async def cmd_helpers(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if not allowed(update):
        return
    if not is_owner(update.effective_user.id):
        await update.message.reply_text("Yordamchilarni faqat bosh admin boshqaradi.")
        return
    if not helpers:
        await update.message.reply_text(
            "Hozircha yordamchi yo'q.\n\nQo'shish uchun: yordamchi botni ochib /start yozsin — "
            "sizga \"✅ Ruxsat berish\" tugmasi keladi."
        )
        return
    lines = ["👥 Yordamchilar (o'chirish uchun ismini bosing):"]
    lines += [f"• {name} — ID {uid}" for uid, name in sorted(helpers.items())]
    await update.message.reply_text("\n".join(lines), reply_markup=helpers_markup())


async def on_button(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    q = update.callback_query
    if not is_owner(q.from_user.id):
        await q.answer("Faqat bosh admin uchun", show_alert=True)
        return
    action, _, raw = (q.data or "").partition(":")
    if not raw.isdigit():
        await q.answer()
        return
    uid = int(raw)
    if action == "add":
        name = requests.pop(uid, "") or helpers.get(uid, "") or str(uid)
        helpers[uid] = name
        dead.discard(uid)
        state["dirty"] = True
        await q.answer("Qo'shildi")
        await q.edit_message_text(f"✅ {name} yordamchi qilib qo'shildi.\nRo'yxat: /yordamchilar")
        try:
            await ctx.bot.send_message(
                uid,
                "✅ Sizga ruxsat berildi! Endi lid xabarlarini shu yerga tashlashingiz mumkin.\n\n"
                "/dastur — guruh qidiruvi dasturi\n/excel — jadval\n/soni — lidlar soni",
            )
        except Exception as e:
            log.warning("Yordamchiga xabar yuborilmadi: %s", e)
    elif action == "no":
        name = requests.pop(uid, str(uid))
        await q.answer("Rad etildi")
        await q.edit_message_text(f"❌ {name} so'rovi rad etildi.")
    elif action == "del":
        name = helpers.pop(uid, None)
        last_pinned.pop(uid, None)
        state["dirty"] = True
        await q.answer("O'chirildi")
        if name is not None:
            try:
                await ctx.bot.send_message(uid, "ℹ️ Botdan foydalanish ruxsatingiz bekor qilindi.")
            except Exception:
                pass
        rest = "\n".join(f"• {n} — ID {i}" for i, n in sorted(helpers.items())) or "Yordamchi qolmadi."
        await q.edit_message_text(f"🗑 {name or uid} o'chirildi.\n\n{rest}", reply_markup=helpers_markup())
    else:
        await q.answer()


async def cmd_excel(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if allowed(update):
        await send_excel(update.effective_chat.id, ctx.bot)


async def cmd_count(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if not allowed(update):
        return
    with lock:
        counts = {h: sum(1 for r in rows if r.get("holat", "Yangi") == h) for h in HOLATLAR}
    lines = [f"Jami: {len(rows)} ta lid"] + [f"• {h}: {n}" for h, n in counts.items() if n]
    await update.message.reply_text("\n".join(lines))


async def cmd_clear(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if not allowed(update):
        return
    if not is_owner(update.effective_user.id):
        await update.message.reply_text("Hamma lidlarni faqat bosh admin o'chira oladi.")
        return
    if ctx.args and ctx.args[0].lower() == "ha":
        with lock:
            rows.clear()
        try:
            os.remove(LOCAL_FILE)
        except OSError:
            pass
        try:
            await ctx.bot.unpin_all_chat_messages(update.effective_chat.id)
        except Exception:
            pass
        last_pinned.pop(update.effective_chat.id, None)
        state["dirty"] = True  # yordamchilarning Excel'i ham bo'shatiladi
        await update.message.reply_text("🗑 Hamma lidlar o'chirildi. Endi lidlarni qaytadan tashlashingiz mumkin.")
    else:
        await update.message.reply_text(
            f"Rostdan ham {len(rows)} ta lidni o'chiraymi? Tasdiqlash uchun yozing:\n/tozalash ha"
        )


async def on_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if not allowed(update):
        return
    text = update.message.text or update.message.caption or ""
    leads = parse_message(text)
    if not leads:
        await update.message.reply_text("Bu xabardan lid topolmadim (ism yoki telefon yo'q).")
        return
    added, dup = add_leads(leads)
    parts = []
    if added:
        names = ", ".join(str(l.get("ism", "?")) for l in rows[-added:])
        parts.append(f"✅ Qo'shildi: {added} ta ({names}) — jami {len(rows)}")
        link = dastur_link()
        if link:
            parts.append(f"📋 Guruhga joylash: {link}")
    if dup:
        parts.append(f"↩️ Takror, o'tkazib yuborildi: {dup} ta")
    await update.message.reply_text("\n".join(parts), disable_web_page_preview=True)
    if added:
        state["dirty"] = True
        schedule_send(update.effective_chat.id, ctx.bot)


async def on_document(update: Update, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    if not allowed(update):
        return
    doc = update.message.document
    if not (doc.file_name or "").lower().endswith(".xlsx"):
        return
    f = await doc.get_file()
    data = bytes(await f.download_as_bytearray())
    try:
        added, dup = add_leads(load_workbook_bytes(data))
    except Exception as e:
        await update.message.reply_text(f"Faylni o'qib bo'lmadi: {e}")
        return
    await update.message.reply_text(f"📥 Fayldan qo'shildi: {added} ta, takror: {dup} ta. Jami: {len(rows)}")
    if added:
        state["dirty"] = True
        schedule_send(update.effective_chat.id, ctx.bot)


async def find_pinned(app: Application, chat_id: int):
    try:
        chat = await app.bot.get_chat(chat_id)
        pm = chat.pinned_message
        if pm and pm.document and (pm.document.file_name or "").endswith(".xlsx"):
            last_pinned[chat_id] = pm.message_id
            return pm
    except Exception as e:
        log.warning("Pin topilmadi (%s): %s", chat_id, e)
    return None


async def restore(app: Application) -> None:
    """Ishga tushganda: avval lokal fayl, bo'lmasa eng so'nggi pin qilingan Excel'dan tiklaydi."""
    asyncio.create_task(persist_loop(app))
    if os.path.exists(LOCAL_FILE):
        with open(LOCAL_FILE, "rb") as f:
            add_leads(load_workbook_bytes(f.read()))
        log.info("Lokal fayldan tiklandi: %d", len(rows))
    pins = [p for p in [await find_pinned(app, a) for a in ADMIN_IDS] if p]
    if not rows and pins:
        latest = max(pins, key=lambda p: p.edit_date or p.date)
        try:
            f = await latest.document.get_file()
            add_leads(load_workbook_bytes(bytes(await f.download_as_bytearray())))
            log.info("Pin qilingan fayldan tiklandi: %d lid, %d yordamchi", len(rows), len(helpers))
        except Exception as e:
            log.warning("Tiklashda xato: %s", e)
    for h in list(helpers):  # yordamchilarning pin qilingan Excel'ini ham topib qo'yamiz
        await find_pinned(app, h)


def main():
    if not BOT_TOKEN:
        raise SystemExit("BOT_TOKEN qo'yilmagan")
    threading.Thread(target=run_web, daemon=True).start()
    if SELF_URL:
        threading.Thread(target=self_ping, daemon=True).start()

    app = (
        Application.builder()
        .token(BOT_TOKEN)
        .connect_timeout(30)
        .read_timeout(30)
        .write_timeout(30)
        .get_updates_connect_timeout(30)
        .get_updates_read_timeout(30)
        .post_init(restore)
        .build()
    )
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("dastur", cmd_dastur))
    app.add_handler(CommandHandler("excel", cmd_excel))
    app.add_handler(CommandHandler("soni", cmd_count))
    app.add_handler(CommandHandler("tozalash", cmd_clear))
    app.add_handler(CommandHandler("yordamchilar", cmd_helpers))
    app.add_handler(CallbackQueryHandler(on_button))
    app.add_handler(MessageHandler(filters.Document.ALL, on_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, on_text))
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
